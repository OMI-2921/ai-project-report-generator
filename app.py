
import io
import os
import re
import textwrap
from datetime import datetime
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

import pandas as pd
import requests
import streamlit as st
from google import genai
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
    Image as PDFImage,
    Table,
    TableStyle,
    KeepTogether,
)


# ---------------------------------------------------------
# APP CONFIGURATION
# ---------------------------------------------------------

st.set_page_config(
    page_title="AI Project Report Generator",
    page_icon="📘",
    layout="wide",
)

APP_TITLE = "AI Project Report Generator"

DEFAULT_CHAPTERS = [
    "Abstract",
    "Introduction",
    "Problem Statement and Objectives",
    "Literature Review",
    "Methodology",
    "System Design",
    "Implementation",
    "Testing and Evaluation",
    "Results and Discussion",
    "Conclusion and Future Scope",
]

WORD_LIMITS = {
    "Concise": 250,
    "Balanced": 330,
    "Detailed": 400,
}

STYLE_OPTIONS = ["Compact", "Standard", "Detailed"]

st.title("📘 AI Project Report Generator")
st.write(
    "Create a structured project report with references, diagrams, "
    "optional data charts, and PDF, Word, and text exports."
)


# ---------------------------------------------------------
# RESEARCH SOURCE SEARCH
# ---------------------------------------------------------

def search_openalex(query, max_results=6):
    """Search OpenAlex for potentially relevant scholarly works."""
    if not query.strip():
        return []

    url = "https://api.openalex.org/works"
    params = {
        "search": query,
        "per-page": max_results,
        "select": (
            "id,doi,title,publication_year,authorships,"
            "primary_location,type"
        ),
    }

    try:
        response = requests.get(url, params=params, timeout=15)
        response.raise_for_status()
        results = response.json().get("results", [])
    except (requests.RequestException, ValueError):
        return []

    sources = []
    for item in results:
        title = (item.get("title") or "").strip()
        if not title:
            continue

        authors = []
        for authorship in item.get("authorships") or []:
            author = (
                authorship.get("author", {}).get("display_name") or ""
            ).strip()
            if author:
                authors.append(author)

        location = item.get("primary_location") or {}
        source = location.get("source") or {}
        journal = (source.get("display_name") or "").strip()

        sources.append({
            "title": title,
            "authors": authors,
            "year": item.get("publication_year"),
            "journal": journal,
            "doi": item.get("doi"),
            "url": item.get("doi") or item.get("id") or "",
        })

    return sources


def search_crossref(query, max_results=6):
    """Fallback scholarly search using Crossref."""
    if not query.strip():
        return []

    try:
        response = requests.get(
            "https://api.crossref.org/works",
            params={"query.bibliographic": query, "rows": max_results},
            timeout=15,
            headers={"User-Agent": "AIProjectReportGenerator/1.0"},
        )
        response.raise_for_status()
        items = response.json().get("message", {}).get("items", [])
    except (requests.RequestException, ValueError):
        return []

    sources = []
    for item in items:
        title_list = item.get("title") or []
        title = title_list[0].strip() if title_list else ""
        if not title:
            continue

        authors = []
        for author in item.get("author") or []:
            given = (author.get("given") or "").strip()
            family = (author.get("family") or "").strip()
            name = " ".join(part for part in [given, family] if part)
            if name:
                authors.append(name)

        date_parts = (
            item.get("published-print", {}).get("date-parts")
            or item.get("published-online", {}).get("date-parts")
            or item.get("created", {}).get("date-parts")
            or []
        )
        year = date_parts[0][0] if date_parts and date_parts[0] else None

        doi = item.get("DOI")
        url = f"https://doi.org/{doi}" if doi else item.get("URL", "")
        container = item.get("container-title") or []

        sources.append({
            "title": title,
            "authors": authors,
            "year": year,
            "journal": container[0] if container else "",
            "doi": doi,
            "url": url,
        })

    return sources


def search_references(query, max_results=8):
    """Search OpenAlex first, then use Crossref as a fallback."""
    sources = search_openalex(query, max_results=max_results)

    if len(sources) < 3:
        crossref_sources = search_crossref(
            query, max_results=max_results
        )
        existing_titles = {
            normalize_title(source["title"]) for source in sources
        }

        for source in crossref_sources:
            key = normalize_title(source["title"])
            if key and key not in existing_titles:
                sources.append(source)
                existing_titles.add(key)

    return sources[:max_results]


def normalize_title(title):
    return re.sub(r"\W+", "", (title or "").lower())


def format_reference(source):
    """Create a readable APA-like reference from available metadata."""
    authors = source.get("authors") or []
    year = source.get("year") or "n.d."
    title = (source.get("title") or "Untitled work").strip()
    journal = (source.get("journal") or "").strip()
    doi = (source.get("doi") or "").strip()
    url = (source.get("url") or "").strip()

    if not authors:
        author_text = "Author metadata unavailable"
    elif len(authors) <= 6:
        author_text = ", ".join(authors)
    else:
        author_text = ", ".join(authors[:6]) + ", et al."

    reference = f"{author_text} ({year}). {title}."

    if journal:
        reference += f" *{journal}*."

    if doi:
        if doi.startswith("http"):
            reference += f" {doi}"
        else:
            reference += f" https://doi.org/{doi}"
    elif url:
        reference += f" {url}"

    return reference


# ---------------------------------------------------------
# DIAGRAM GENERATION
# ---------------------------------------------------------

def figure_to_png(fig):
    buffer = io.BytesIO()
    fig.savefig(
        buffer,
        format="png",
        dpi=180,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(fig)
    buffer.seek(0)
    return buffer.getvalue()


def add_box(ax, x, y, width, height, label, fontsize=10):
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.03,rounding_size=0.04",
        linewidth=1.2,
        edgecolor="#334155",
        facecolor="#EFF6FF",
    )
    ax.add_patch(patch)
    ax.text(
        x + width / 2,
        y + height / 2,
        label,
        ha="center",
        va="center",
        fontsize=fontsize,
        wrap=True,
        color="#172554",
    )


def add_arrow(ax, start, end):
    arrow = FancyArrowPatch(
        start,
        end,
        arrowstyle="-|>",
        mutation_scale=14,
        linewidth=1.3,
        color="#475569",
    )
    ax.add_patch(arrow)


def make_workflow_diagram(title):
    fig, ax = plt.subplots(figsize=(10, 2.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2.3)
    ax.axis("off")

    ax.text(
        5, 2.05, f"Workflow: {title}",
        ha="center", va="center", fontsize=13, weight="bold",
    )

    labels = [
        "Input",
        "Processing",
        "Validation",
        "Output",
    ]
    xs = [0.4, 2.85, 5.3, 7.75]

    for x, label in zip(xs, labels):
        add_box(ax, x, 0.75, 1.65, 0.7, label)

    for x in [2.1, 4.55, 7.0]:
        add_arrow(ax, (x, 1.1), (x + 0.55, 1.1))

    return figure_to_png(fig)


def make_architecture_diagram(title):
    fig, ax = plt.subplots(figsize=(9, 4))
    ax.set_xlim(0, 9)
    ax.set_ylim(0, 4)
    ax.axis("off")

    ax.text(
        4.5, 3.65, f"Conceptual System Architecture: {title}",
        ha="center", va="center", fontsize=12, weight="bold",
    )

    add_box(ax, 0.5, 2.1, 2.0, 0.8, "User Interface")
    add_box(ax, 3.5, 2.1, 2.0, 0.8, "Application Logic")
    add_box(ax, 6.5, 2.1, 2.0, 0.8, "Output / Services")
    add_arrow(ax, (2.5, 2.5), (3.5, 2.5))
    add_arrow(ax, (5.5, 2.5), (6.5, 2.5))

    add_box(ax, 3.5, 0.65, 2.0, 0.7, "Data / Storage")
    add_arrow(ax, (4.5, 2.1), (4.5, 1.35))

    ax.text(
        4.5, 0.25,
        "Conceptual illustration — adapt to the actual implementation.",
        ha="center", va="center", fontsize=8, color="#475569",
    )

    return figure_to_png(fig)


def make_timeline_diagram(title):
    phases = [
        "Planning",
        "Design",
        "Development",
        "Testing",
        "Documentation",
    ]
    starts = [0, 1, 2, 3, 4]
    durations = [1, 1, 2, 1, 1]

    fig, ax = plt.subplots(figsize=(9, 3.4))
    ax.barh(phases, durations, left=starts)
    ax.set_title(f"Illustrative Project Timeline: {title}")
    ax.set_xlabel("Relative project phase (not actual dates)")
    ax.set_ylabel("Phase")
    ax.grid(axis="x", alpha=0.25)
    ax.invert_yaxis()

    fig.text(
        0.5, 0.01,
        "Illustrative plan only; replace with the project's real schedule.",
        ha="center", fontsize=8,
    )
    fig.tight_layout(rect=[0, 0.06, 1, 1])
    return figure_to_png(fig)


def make_entity_diagram(title):
    fig, ax = plt.subplots(figsize=(8, 3.5))
    ax.set_xlim(0, 8)
    ax.set_ylim(0, 3.5)
    ax.axis("off")

    ax.text(
        4, 3.2, f"Conceptual Data Model: {title}",
        ha="center", va="center", fontsize=12, weight="bold",
    )

    add_box(ax, 0.5, 1.3, 2.0, 1.0, "Entity A\nID\nAttributes")
    add_box(ax, 5.5, 1.3, 2.0, 1.0, "Entity B\nID\nAttributes")
    add_arrow(ax, (2.5, 1.8), (5.5, 1.8))

    ax.text(
        4, 2.0, "relationship",
        ha="center", va="bottom", fontsize=8,
    )
    ax.text(
        4, 0.65,
        "Generic example — customize entities and relationships for your project.",
        ha="center", fontsize=8,
    )

    return figure_to_png(fig)


def make_csv_chart(dataframe, x_column, y_column, chart_type, title):
    """Build a chart using columns from the user's uploaded CSV."""
    df = dataframe[[x_column, y_column]].dropna().copy()

    if df.empty:
        raise ValueError("There are no complete rows for the selected columns.")

    if chart_type == "Line":
        fig, ax = plt.subplots(figsize=(8, 4.2))
        ax.plot(df[x_column].astype(str), pd.to_numeric(
            df[y_column], errors="coerce"
        ), marker="o")
        ax.set_xlabel(x_column)
        ax.set_ylabel(y_column)
        ax.tick_params(axis="x", rotation=35)

    elif chart_type == "Bar":
        fig, ax = plt.subplots(figsize=(8, 4.2))
        values = pd.to_numeric(df[y_column], errors="coerce")
        ax.bar(df[x_column].astype(str), values)
        ax.set_xlabel(x_column)
        ax.set_ylabel(y_column)
        ax.tick_params(axis="x", rotation=35)

    else:
        fig, ax = plt.subplots(figsize=(8, 4.2))
        values = pd.to_numeric(df[y_column], errors="coerce")
        ax.scatter(df[x_column].astype(str), values)
        ax.set_xlabel(x_column)
        ax.set_ylabel(y_column)
        ax.tick_params(axis="x", rotation=35)

    ax.set_title(title)
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    return figure_to_png(fig)


# ---------------------------------------------------------
# GEMINI REPORT GENERATION
# ---------------------------------------------------------

def build_report_prompt(
    project_title,
    project_type,
    description,
    audience,
    selected_chapters,
    target_pages,
    formatting_style,
    content_density,
    sources,
    data_description,
):
    words_per_page = {
        "Compact": 420,
        "Standard": 330,
        "Detailed": 260,
    }[formatting_style]

    density_factor = {
        "Concise": 0.8,
        "Balanced": 1.0,
        "Detailed": 1.15,
    }[content_density]

    estimated_words = int(
        max(500, target_pages * words_per_page * density_factor)
    )

    # One model response may not be able to produce very long reports.
    # Keep the requested budget bounded to reduce truncation risk.
    requested_words = min(estimated_words, 10000)

    source_text = "\n".join(
        f"- {format_reference(source)}"
        for source in sources
    ) or "- No verified scholarly sources were retrieved."

    chapter_instructions = "\n".join(
        f"{index}. {chapter}"
        for index, chapter in enumerate(selected_chapters, start=1)
    )

    return f"""
You are an academic project-report writing assistant.

PROJECT DETAILS
Title: {project_title}
Project type: {project_type}
Intended audience: {audience}
Project description and requirements:
{description}

REPORT SETTINGS
Target PDF length: approximately {target_pages} pages.
Formatting style: {formatting_style}
Content density: {content_density}
Approximate content word budget: {requested_words} words.
Selected chapters:
{chapter_instructions}

DATA PROVIDED BY USER
{data_description or "No dataset description supplied."}

POTENTIALLY RELEVANT SOURCES RETRIEVED FROM SCHOLARLY SEARCH
{source_text}

CRITICAL ACCURACY RULES
1. Do not invent experiments, participants, measurements, survey responses,
   numerical findings, performance scores, or evaluation outcomes.
2. If the user has not supplied results, clearly state that results are
   pending, proposed, or illustrative rather than claiming they occurred.
3. Do not invent references, authors, article titles, journals, DOIs, or URLs.
4. Use the supplied source list only when it genuinely supports a claim.
   If a source is not relevant, do not cite it.
5. Do not claim that you accessed or reviewed a full paper when only its
   metadata is available.
6. Distinguish planned functionality from implemented functionality.
7. Do not claim that a system was tested or deployed unless the user supplied
   evidence that this happened.
8. Use clear, formal academic language suitable for a student project report.
9. Avoid repetition, filler, fake quotations, and unsupported statistics.
10. Explain assumptions and limitations openly.

WRITING INSTRUCTIONS
- Write a useful, coherent report with substantial content under every
  selected chapter, proportional to the requested length.
- Explain concepts, design decisions, procedures, and limitations in detail.
- Use numbered or bulleted lists where they improve readability.
- Use Markdown subheadings inside chapters when helpful.
- Do not create a References chapter. References will be added separately
  from retrieved metadata.
- If the project details are insufficient for a factual claim, describe what
  needs to be confirmed rather than guessing.
- The target page count is an estimate, not a guarantee. The final PDF may
  have a different page count because of pagination and diagrams.

OUTPUT FORMAT
Return only the report chapters, using exactly these markers to start each
chapter. Do not omit or rename markers:

[[CHAPTER: Chapter Name]]

Write that chapter's content here.

Repeat the marker format for every selected chapter, preserving the chapter
names and order exactly. Do not add an introduction outside the markers.
""".strip()


def parse_generated_chapters(response_text, selected_chapters):
    """Parse model output into chapters using explicit chapter markers."""
    text = (response_text or "").strip()
    pattern = re.compile(r"\[\[CHAPTER:\s*(.*?)\s*\]\]", re.IGNORECASE)
    matches = list(pattern.finditer(text))

    chapters = {}

    if matches:
        for index, match in enumerate(matches):
            name = match.group(1).strip()
            start = match.end()
            end = (
                matches[index + 1].start()
                if index + 1 < len(matches)
                else len(text)
            )
            content = text[start:end].strip()

            if name and content:
                chapters[name] = content

    # Match headings case-insensitively to the user-selected chapter names.
    normalized = {
        re.sub(r"\s+", " ", name.strip().lower()): content
        for name, content in chapters.items()
    }

    result = {}
    for chapter in selected_chapters:
        key = re.sub(r"\s+", " ", chapter.strip().lower())
        result[chapter] = normalized.get(key, "")

    # If marker parsing failed completely, preserve the response rather than
    # silently discarding the model's content.
    if not any(result.values()) and text:
        first_chapter = selected_chapters[0]
        result[first_chapter] = text

    return result


def generate_report(
    api_key,
    model_name,
    project_title,
    project_type,
    description,
    audience,
    selected_chapters,
    target_pages,
    formatting_style,
    content_density,
    sources,
    data_description,
):
    client = genai.Client(api_key=api_key)

    prompt = build_report_prompt(
        project_title=project_title,
        project_type=project_type,
        description=description,
        audience=audience,
        selected_chapters=selected_chapters,
        target_pages=target_pages,
        formatting_style=formatting_style,
        content_density=content_density,
        sources=sources,
        data_description=data_description,
    )

    response = client.models.generate_content(
        model=model_name,
        contents=prompt,
    )

    response_text = getattr(response, "text", None)
    if not response_text:
        raise RuntimeError(
            "The AI returned an empty response. Please try again."
        )

    return parse_generated_chapters(response_text, selected_chapters)


# ---------------------------------------------------------
# DOCUMENT EXPORT HELPERS
# ---------------------------------------------------------

def clean_markdown(text):
    """Lightweight Markdown cleanup for document exports."""
    text = re.sub(r"```.*?```", "[Code block omitted in document export]", text,
                  flags=re.DOTALL)
    text = re.sub(r"!\[.*?\]\(.*?\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"(?m)^\s*[-*]\s+", "• ", text)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)", r"\1", text)
    return text.strip()


def split_markdown_blocks(text):
    """Yield basic heading, paragraph, and list blocks."""
    lines = (text or "").splitlines()
    paragraph_lines = []
    list_lines = []

    def flush_paragraph():
        nonlocal paragraph_lines
        if paragraph_lines:
            value = " ".join(line.strip() for line in paragraph_lines).strip()
            if value:
                yield ("paragraph", value)
            paragraph_lines = []

    def flush_list():
        nonlocal list_lines
        if list_lines:
            yield ("list", list_lines[:])
            list_lines = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            yield from flush_paragraph()
            yield from flush_list()
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            yield from flush_paragraph()
            yield from flush_list()
            yield ("heading", (len(heading_match.group(1)), heading_match.group(2)))
            continue

        if re.match(r"^[-*]\s+", stripped):
            yield from flush_paragraph()
            list_lines.append(re.sub(r"^[-*]\s+", "", stripped))
            continue

        yield from flush_list()
        paragraph_lines.append(stripped)

    yield from flush_paragraph()
    yield from flush_list()


def add_chapter_to_word(doc, title, content):
    doc.add_heading(title, level=1)

    for kind, value in split_markdown_blocks(content):
        if kind == "heading":
            level, heading = value
            doc.add_heading(heading, level=min(level + 1, 4))

        elif kind == "paragraph":
            paragraph = doc.add_paragraph(clean_markdown(value))
            paragraph.paragraph_format.space_after = Pt(6)

        elif kind == "list":
            for item in value:
                doc.add_paragraph(clean_markdown(item), style="List Bullet")


def create_word_report(title, chapters, references, visuals):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = 1
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(24)

    subtitle = doc.add_paragraph()
    subtitle.alignment = 1
    subtitle.add_run("Project Report").italic = True

    date_paragraph = doc.add_paragraph()
    date_paragraph.alignment = 1
    date_paragraph.add_run(
        f"Generated on {datetime.now().strftime('%d %B %Y')}"
    )

    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    for chapter_name in chapters:
        doc.add_paragraph(chapter_name, style="List Bullet")
    if visuals:
        doc.add_paragraph("Figures and diagrams", style="List Bullet")
    if references:
        doc.add_paragraph("References", style="List Bullet")

    doc.add_page_break()

    visual_index = 0
    for chapter_name, content in chapters.items():
        if not content.strip():
            continue

        add_chapter_to_word(doc, chapter_name, content)

        # Insert appropriate conceptual visuals near design/implementation.
        lower_name = chapter_name.lower()
        if visuals and (
            "design" in lower_name
            or "methodology" in lower_name
            or "implementation" in lower_name
        ):
            for visual in visuals:
                if visual_index >= len(visuals):
                    break
                if visual.get("inserted"):
                    continue

                doc.add_paragraph(visual["caption"])
                doc.add_picture(
                    io.BytesIO(visual["image"]),
                    width=Inches(6.0),
                )
                visual["inserted"] = True
                visual_index += 1

    remaining_visuals = [
        visual for visual in visuals if not visual.get("inserted")
    ]
    if remaining_visuals:
        doc.add_heading("Figures and Diagrams", level=1)
        for visual in remaining_visuals:
            doc.add_paragraph(visual["caption"])
            doc.add_picture(
                io.BytesIO(visual["image"]),
                width=Inches(6.0),
            )

    if references:
        doc.add_heading("References", level=1)
        for reference in references:
            doc.add_paragraph(reference)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    for visual in visuals:
        visual.pop("inserted", None)

    return buffer.getvalue()


def reportlab_image(image_bytes, max_width=6.5 * inch, max_height=4.0 * inch):
    image = PDFImage(io.BytesIO(image_bytes))
    ratio = min(
        max_width / image.imageWidth,
        max_height / image.imageHeight,
    )
    image.drawWidth = image.imageWidth * ratio
    image.drawHeight = image.imageHeight * ratio
    return image


def add_pdf_page_number(canvas, document, page_counter):
    page_counter["count"] = canvas.getPageNumber()
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.grey)
    canvas.drawCentredString(
        A4[0] / 2,
        0.35 * inch,
        f"Page {canvas.getPageNumber()}",
    )
    canvas.restoreState()


def create_pdf_report(title, chapters, references, visuals):
    buffer = io.BytesIO()
    page_counter = {"count": 0}

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.65 * inch,
        title=title,
        author="AI Project Report Generator",
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=23,
        leading=29,
        alignment=TA_CENTER,
        spaceAfter=20,
    ))
    styles.add(ParagraphStyle(
        name="ChapterHeading",
        parent=styles["Heading1"],
        fontSize=16,
        leading=20,
        spaceBefore=10,
        spaceAfter=10,
        keepWithNext=True,
    ))
    styles.add(ParagraphStyle(
        name="SubHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        spaceBefore=8,
        spaceAfter=5,
        keepWithNext=True,
    ))
    styles.add(ParagraphStyle(
        name="ReportBody",
        parent=styles["BodyText"],
        fontSize=9.5,
        leading=13.2,
        spaceAfter=7,
    ))
    styles.add(ParagraphStyle(
        name="ReferenceBody",
        parent=styles["BodyText"],
        fontSize=8.5,
        leading=11,
        leftIndent=12,
        firstLineIndent=-12,
        spaceAfter=7,
    ))
    styles.add(ParagraphStyle(
        name="SmallCaption",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
        alignment=TA_CENTER,
        textColor=colors.darkgrey,
        spaceAfter=6,
    ))

    story = [
        Spacer(1, 1.3 * inch),
        Paragraph(escape(title), styles["ReportTitle"]),
        Paragraph("Project Report", styles["Heading2"]),
        Spacer(1, 0.25 * inch),
        Paragraph(
            f"Generated on {datetime.now().strftime('%d %B %Y')}",
            styles["BodyText"],
        ),
        PageBreak(),
        Paragraph("Contents", styles["ChapterHeading"]),
    ]

    for chapter_name in chapters:
        story.append(Paragraph(
            escape(chapter_name), styles["ReportBody"]
        ))
    if visuals:
        story.append(Paragraph("Figures and diagrams", styles["ReportBody"]))
    if references:
        story.append(Paragraph("References", styles["ReportBody"]))

    story.append(PageBreak())

    for chapter_name, content in chapters.items():
        if not content.strip():
            continue

        story.append(Paragraph(
            escape(chapter_name), styles["ChapterHeading"]
        ))

        for kind, value in split_markdown_blocks(content):
            if kind == "heading":
                level, heading = value
                style = (
                    styles["SubHeading"]
                    if level <= 2
                    else styles["Heading3"]
                )
                story.append(Paragraph(escape(heading), style))

            elif kind == "paragraph":
                paragraph = clean_markdown(value)
                story.append(Paragraph(escape(paragraph), styles["ReportBody"]))

            elif kind == "list":
                for item in value:
                    story.append(Paragraph(
                        "• " + escape(clean_markdown(item)),
                        styles["ReportBody"],
                    ))

        lower_name = chapter_name.lower()
        if visuals and (
            "design" in lower_name
            or "methodology" in lower_name
            or "implementation" in lower_name
        ):
            for visual in visuals:
                if visual.get("inserted"):
                    continue
                story.append(Spacer(1, 6))
                story.append(Paragraph(
                    escape(visual["caption"]), styles["SmallCaption"]
                ))
                story.append(reportlab_image(visual["image"]))
                story.append(Spacer(1, 10))
                visual["inserted"] = True

    remaining_visuals = [
        visual for visual in visuals if not visual.get("inserted")
    ]
    if remaining_visuals:
        story.append(Paragraph(
            "Figures and Diagrams", styles["ChapterHeading"]
        ))
        for visual in remaining_visuals:
            story.append(Paragraph(
                escape(visual["caption"]), styles["SmallCaption"]
            ))
            story.append(reportlab_image(visual["image"]))
            story.append(Spacer(1, 10))

    if references:
        story.append(Paragraph("References", styles["ChapterHeading"]))
        for reference in references:
            story.append(Paragraph(
                escape(reference), styles["ReferenceBody"]
            ))

    document.build(
        story,
        onFirstPage=lambda canvas, doc: add_pdf_page_number(
            canvas, doc, page_counter
        ),
        onLaterPages=lambda canvas, doc: add_pdf_page_number(
            canvas, doc, page_counter
        ),
    )

    pdf_bytes = buffer.getvalue()

    for visual in visuals:
        visual.pop("inserted", None)

    return pdf_bytes, page_counter["count"]


def create_text_report(title, chapters, references):
    sections = [title, "=" * len(title), ""]

    for chapter_name, content in chapters.items():
        if content.strip():
            sections.extend([
                chapter_name,
                "-" * len(chapter_name),
                clean_markdown(content),
                "",
            ])

    if references:
        sections.extend(["References", "----------"])
        sections.extend(references)

    return "\n".join(sections)


# ---------------------------------------------------------
# SIDEBAR SETTINGS AND INPUT FORM
# ---------------------------------------------------------

with st.sidebar:
    st.header("Report settings")

    target_pages = st.slider(
        "Target PDF pages",
        min_value=5,
        max_value=100,
        value=15,
        step=5,
        help=(
            "This guides the report length. Actual pages depend on the "
            "generated text, figures, and PDF layout."
        ),
    )

    formatting_style = st.selectbox(
        "Formatting style",
        STYLE_OPTIONS,
        index=1,
    )

    content_density = st.selectbox(
        "Content density",
        ["Concise", "Balanced", "Detailed"],
        index=1,
    )

    st.divider()
    st.subheader("Include visuals")

    include_workflow = st.checkbox("Workflow diagram", value=True)
    include_architecture = st.checkbox(
        "System architecture diagram", value=True
    )
    include_timeline = st.checkbox(
        "Illustrative project timeline", value=True
    )
    include_entity_model = st.checkbox(
        "Conceptual data model", value=False
    )


st.subheader("1. Project information")

project_title = st.text_input(
    "Project title",
    placeholder="e.g. AI-Based Student Performance Prediction",
)

col1, col2 = st.columns(2)

with col1:
    project_type = st.selectbox(
        "Project type",
        [
            "Software / Web Application",
            "Artificial Intelligence / Machine Learning",
            "Data Science / Analytics",
            "Research Project",
            "Mobile Application",
            "IoT / Embedded System",
            "Business / Management",
            "Other",
        ],
    )

with col2:
    audience = st.selectbox(
        "Report audience",
        [
            "College / University",
            "Technical Review",
            "Business / Management",
            "General Audience",
        ],
    )

description = st.text_area(
    "Project description and requirements",
    placeholder=(
        "Explain the problem, goals, planned or completed features, "
        "technologies, methods, and any real findings you have."
    ),
    height=170,
)

st.subheader("2. Choose report chapters")

selected_chapters = st.multiselect(
    "Select the chapters to generate",
    options=DEFAULT_CHAPTERS,
    default=DEFAULT_CHAPTERS,
)

st.subheader("3. Optional data for charts")

uploaded_csv = st.file_uploader(
    "Upload a CSV file to create a chart from your data",
    type=["csv"],
    help=(
        "Charts made from this file use its actual rows and selected columns. "
        "Do not upload sensitive personal information."
    ),
)

csv_dataframe = None
chart_image = None
chart_caption = ""
data_description = ""

if uploaded_csv is not None:
    try:
        csv_dataframe = pd.read_csv(uploaded_csv)

        if csv_dataframe.empty:
            st.warning("The uploaded CSV has no data rows.")
            csv_dataframe = None
        else:
            st.write("Preview of your uploaded data:")
            st.dataframe(csv_dataframe.head(10), use_container_width=True)

            columns = list(csv_dataframe.columns)
            chart_col1, chart_col2 = st.columns(2)

            with chart_col1:
                x_column = st.selectbox(
                    "Chart category / X-axis",
                    options=columns,
                    key="chart_x",
                )

            numeric_columns = [
                column for column in columns
                if pd.api.types.is_numeric_dtype(csv_dataframe[column])
            ]

            with chart_col2:
                if numeric_columns:
                    y_column = st.selectbox(
                        "Numeric Y-axis",
                        options=numeric_columns,
                        key="chart_y",
                    )
                else:
                    y_column = None
                    st.info(
                        "No numeric columns were detected. Add a numeric "
                        "column to create a chart."
                    )

            chart_type = st.selectbox(
                "Chart type",
                ["Bar", "Line", "Scatter"],
                key="chart_type",
            )

            chart_title = st.text_input(
                "Chart title",
                value="Chart generated from uploaded data",
            )

            if y_column:
                try:
                    chart_image = make_csv_chart(
                        csv_dataframe,
                        x_column,
                        y_column,
                        chart_type,
                        chart_title,
                    )
                    chart_caption = (
                        f"Figure: {chart_title}. Generated from the user's "
                        f"uploaded CSV using {x_column} and {y_column}."
                    )
                    st.image(chart_image, caption=chart_caption)
                    data_description = (
                        f"The user uploaded a CSV with {len(csv_dataframe)} "
                        f"rows and {len(csv_dataframe.columns)} columns. "
                        f"The selected chart uses X='{x_column}', "
                        f"Y='{y_column}', type='{chart_type}'. "
                        "Do not infer findings that have not been calculated "
                        "or supported by the data."
                    )
                except (ValueError, TypeError) as error:
                    st.warning(f"Could not create the chart: {error}")

    except (ValueError, UnicodeDecodeError, pd.errors.ParserError) as error:
        st.error(f"Could not read this CSV file: {error}")
        csv_dataframe = None


# ---------------------------------------------------------
# GENERATE REPORT
# ---------------------------------------------------------

st.divider()

if target_pages >= 40:
    st.warning(
        "Long reports may exceed the AI model's single-response output limit. "
        "The requested page count is a target, not a guarantee. If the report "
        "is much shorter than expected, try generating a smaller report first."
    )

if st.button("Generate project report", type="primary", use_container_width=True):
    if not project_title.strip():
        st.error("Please enter a project title.")
        st.stop()

    if not description.strip():
        st.error("Please describe your project so the report can be tailored.")
        st.stop()

    if not selected_chapters:
        st.error("Please select at least one chapter.")
        st.stop()

    try:
        api_key = st.secrets.get("GEMINI_API_KEY", "")
        model_name = st.secrets.get(
            "GEMINI_MODEL",
            "gemini-2.5-flash",
        )
    except Exception:
        api_key = os.getenv("GEMINI_API_KEY", "")
        model_name = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

    if not api_key:
        st.error(
            "Gemini API key not found. Add GEMINI_API_KEY to your "
            "Streamlit app's Secrets settings."
        )
        st.stop()

    search_query = f"{project_title} {project_type} {description[:250]}"

    with st.status("Preparing your report...", expanded=True) as status:
        st.write("Searching scholarly source metadata...")
        sources = search_references(search_query, max_results=8)

        st.write("Generating report chapters with the AI model...")
        try:
            chapters = generate_report(
                api_key=api_key,
                model_name=model_name,
                project_title=project_title.strip(),
                project_type=project_type,
                description=description.strip(),
                audience=audience,
                selected_chapters=selected_chapters,
                target_pages=target_pages,
                formatting_style=formatting_style,
                content_density=content_density,
                sources=sources,
                data_description=data_description,
            )
        except Exception as error:
            status.update(label="Report generation failed", state="error")
            st.error(
                "The AI could not generate the report. Check your API key, "
                "model name, quota, and connection, then try again."
            )
            st.exception(error)
            st.stop()

        st.write("Preparing diagrams and charts...")
        visuals = []

        try:
            if include_workflow:
                visuals.append({
                    "caption": "Workflow diagram — conceptual overview.",
                    "image": make_workflow_diagram(project_title),
                })

            if include_architecture:
                visuals.append({
                    "caption": (
                        "System architecture — conceptual illustration; "
                        "verify and adapt it to the actual system."
                    ),
                    "image": make_architecture_diagram(project_title),
                })

            if include_timeline:
                visuals.append({
                    "caption": (
                        "Illustrative project timeline — not actual dates "
                        "or evidence of completed work."
                    ),
                    "image": make_timeline_diagram(project_title),
                })

            if include_entity_model:
                visuals.append({
                    "caption": (
                        "Conceptual data model — generic example to customize."
                    ),
                    "image": make_entity_diagram(project_title),
                })

            if chart_image:
                visuals.append({
                    "caption": chart_caption,
                    "image": chart_image,
                })

        except Exception as error:
            st.warning(f"One or more visuals could not be generated: {error}")

        references = [format_reference(source) for source in sources]

        st.write("Creating PDF and Word exports...")
        try:
            pdf_bytes, actual_pdf_pages = create_pdf_report(
                project_title.strip(),
                chapters,
                references,
                visuals,
            )

            word_bytes = create_word_report(
                project_title.strip(),
                chapters,
                references,
                visuals,
            )

            text_report = create_text_report(
                project_title.strip(),
                chapters,
                references,
            )

        except Exception as error:
            status.update(label="Export creation failed", state="error")
            st.error("The report was generated, but an export failed.")
            st.exception(error)
            st.stop()

        status.update(label="Report ready", state="complete")

    st.session_state["generated_report"] = {
        "title": project_title.strip(),
        "chapters": chapters,
        "references": references,
        "pdf": pdf_bytes,
        "pdf_pages": actual_pdf_pages,
        "word": word_bytes,
        "text": text_report,
        "target_pages": target_pages,
    }


# ---------------------------------------------------------
# DISPLAY AND DOWNLOAD GENERATED REPORT
# ---------------------------------------------------------

report = st.session_state.get("generated_report")

if report:
    st.divider()
    st.subheader("4. Your report is ready")

    actual_pages = report["pdf_pages"]
    target = report["target_pages"]

    metric1, metric2, metric3 = st.columns(3)
    metric1.metric("Target PDF pages", target)
    metric2.metric("Actual PDF pages", actual_pages)
    metric3.metric(
        "Difference",
        f"{actual_pages - target:+d} pages",
    )

    if actual_pages == target:
        st.success("The PDF matches your target page count.")
    elif abs(actual_pages - target) <= 2:
        st.info(
            "The PDF is close to your target. Page counts can vary with "
            "font rendering, figures, and paragraph layout."
        )
    else:
        st.warning(
            "The PDF differs from your target. Try changing the content "
            "density or target page count and generating again. A target "
            "page count cannot be guaranteed by the AI model."
        )

    st.write("Download your generated files:")

    col_pdf, col_word, col_text = st.columns(3)

    safe_name = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        report["title"],
    ).strip("_") or "project_report"

    with col_pdf:
        st.download_button(
            "📄 Download PDF",
            data=report["pdf"],
            file_name=f"{safe_name}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    with col_word:
        st.download_button(
            "📝 Download Word",
            data=report["word"],
            file_name=f"{safe_name}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )

    with col_text:
        st.download_button(
            "📃 Download TXT",
            data=report["text"],
            file_name=f"{safe_name}.txt",
            mime="text/plain",
            use_container_width=True,
        )

    with st.expander("Preview generated chapters"):
        for chapter_name, content in report["chapters"].items():
            if content.strip():
                st.markdown(f"### {chapter_name}")
                st.markdown(content)

    with st.expander("Review retrieved references"):
        if report["references"]:
            for index, reference in enumerate(report["references"], start=1):
                st.write(f"{index}. {reference}")
        else:
            st.info(
                "No scholarly references were retrieved. Search for relevant "
                "sources separately and verify bibliographic details before "
                "submitting your report."
            )

st.caption(
    "Always verify citations, technical details, diagrams, and claims before "
    "submitting a report. Conceptual illustrations are not evidence of actual "
    "implementation or experimental results."
)
